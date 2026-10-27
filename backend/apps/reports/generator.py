import os
import csv
from datetime import datetime
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.conf import settings


def generate_pdf_report(report_title: str, date_str: str, summary: dict, visitors: list, camera_activity: list, output_filepath: str):
    doc = SimpleDocTemplate(
        output_filepath,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    styles = getSampleStyleSheet()
    story = []

    # Title Banner
    title_style = ParagraphStyle(
        'ReportTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=20,
        textColor=colors.HexColor('#0F172A'),
        spaceAfter=4
    )
    subtitle_style = ParagraphStyle(
        'ReportSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=11,
        textColor=colors.HexColor('#64748B'),
        spaceAfter=14
    )

    story.append(Paragraph("VISIONTRACK AI", title_style))
    story.append(Paragraph(f"{report_title} &bull; Generated for: {date_str}", subtitle_style))
    story.append(Spacer(1, 10))

    # Summary KPI Table
    kpi_data = [
        [
            f"Unique People\n{summary.get('unique_people', 0)}",
            f"Total Visits\n{summary.get('total_visits', 0)}",
            f"Vehicles\n{summary.get('vehicles', 0)}",
            f"Animals\n{summary.get('animals', 0)}",
            f"Peak Hour\n{summary.get('peak_hour', '10:00 - 11:00')}"
        ]
    ]
    kpi_table = Table(kpi_data, colWidths=[108, 108, 108, 108, 108])
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F1F5F9')),
        ('TEXTCOLOR', (0, 0), (-1, -1), colors.HexColor('#1E293B')),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 10),
        ('TOPPADDING', (0, 0), (-1, -1), 10),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#CBD5E1')),
    ]))
    story.append(kpi_table)
    story.append(Spacer(1, 18))

    # Section Heading
    heading_style = ParagraphStyle(
        'SectionHeading',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=14,
        textColor=colors.HexColor('#1E293B'),
        spaceAfter=8
    )
    story.append(Paragraph("Historical Visitors & Observed Objects", heading_style))

    # Visitor Table
    headers = ["Object ID", "Type", "First Seen", "Last Seen", "Visits", "Duration"]
    table_rows = [headers]

    for v in visitors[:50]:
        table_rows.append([
            v.get("internal_id", "N/A"),
            v.get("object_type", "Person").capitalize(),
            v.get("first_seen", "")[:16].replace("T", " "),
            v.get("last_seen", "")[:16].replace("T", " "),
            str(v.get("total_visits", 1)),
            f"{v.get('total_duration_seconds', 0) // 60}m"
        ])

    vis_table = Table(table_rows, colWidths=[110, 80, 110, 110, 50, 80])
    vis_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E293B')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
    ]))
    story.append(vis_table)
    story.append(Spacer(1, 18))

    # Camera Activity Section
    if camera_activity:
        story.append(Paragraph("Camera Activity Summary", heading_style))
        cam_headers = ["Camera ID", "Camera Name", "Visits Recorded"]
        cam_rows = [cam_headers]
        for cam in camera_activity:
            cam_rows.append([cam.get("camera_id", ""), cam.get("name", ""), f"{cam.get('visits', 0)} visits"])
        cam_table = Table(cam_rows, colWidths=[120, 260, 160])
        cam_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#334155')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
        ]))
        story.append(cam_table)

    doc.build(story)


def generate_excel_report(report_title: str, date_str: str, summary: dict, visitors: list, output_filepath: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Visitor Report"

    # Header
    ws.append(["VISIONTRACK AI", report_title, date_str])
    ws.append([])
    ws.append(["SUMMARY METRICS"])
    ws.append(["Unique People", summary.get("unique_people", 0)])
    ws.append(["Total Visits", summary.get("total_visits", 0)])
    ws.append(["Vehicles", summary.get("vehicles", 0)])
    ws.append(["Animals", summary.get("animals", 0)])
    ws.append([])

    # Table
    ws.append(["Object ID", "Type", "First Seen", "Last Seen", "Visits", "Duration (sec)"])
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for cell in ws[8]:
        cell.fill = header_fill
        cell.font = header_font

    for v in visitors:
        ws.append([
            v.get("internal_id"),
            v.get("object_type"),
            str(v.get("first_seen")),
            str(v.get("last_seen")),
            v.get("total_visits"),
            v.get("total_duration_seconds")
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max(max_len + 3, 12)

    wb.save(output_filepath)


def generate_csv_report(summary: dict, visitors: list, output_filepath: str):
    with open(output_filepath, 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(["VISIONTRACK AI - VISITOR REPORT"])
        writer.writerow(["Unique People", summary.get("unique_people", 0)])
        writer.writerow(["Total Visits", summary.get("total_visits", 0)])
        writer.writerow(["Vehicles", summary.get("vehicles", 0)])
        writer.writerow(["Animals", summary.get("animals", 0)])
        writer.writerow([])
        writer.writerow(["Object ID", "Type", "First Seen", "Last Seen", "Visits", "Duration (sec)"])
        for v in visitors:
            writer.writerow([
                v.get("internal_id"),
                v.get("object_type"),
                v.get("first_seen"),
                v.get("last_seen"),
                v.get("total_visits"),
                v.get("total_duration_seconds")
            ])
